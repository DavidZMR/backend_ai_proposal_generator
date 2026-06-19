"""
docx_builder.py — Construcción de documentos de propuesta en formato DOCX.

Permite generar propuestas comerciales en formato Word aplicando estilos
corporativos de LinkThinks.
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Colores Corporativos LinkThinks
TEAL_HEX = (22, 181, 160)  # #16B5A0
NAVY_HEX = (11, 30, 58)    # #0B1E3A

def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r._r.append(fldChar)
    r._r.append(instrText)
    r._r.append(fldChar2)
    r._r.append(fldChar3)

def build_docx(sections: dict, prospect_name: str, output_path: str) -> str:
    """
    Genera un archivo DOCX con el contenido de la propuesta usando estilos corporativos.
    
    Args:
        sections: Diccionario con el contenido generado de las secciones.
        prospect_name: Nombre del cliente/prospecto.
        output_path: Ruta absoluta donde se guardará el archivo DOCX.
        
    Returns:
        La misma ruta del archivo generado.
    """
    doc = Document()
    
    # 1. Portada
    
    logo_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "storage", "assets", "cover_image.png")
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(logo_path, width=Inches(6))

    # Espaciado superior para centrar verticalmente la portada
    for _ in range(3):
        doc.add_paragraph()
        
    title = doc.add_paragraph("Propuesta de Servicios")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.runs[0]
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(32)
    title_run.font.color.rgb = RGBColor(*NAVY_HEX)
    title_run.font.bold = True
    
    doc.add_paragraph()
    
    subtitle = doc.add_paragraph(f"Para: {prospect_name}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.runs[0]
    subtitle_run.font.name = 'Arial'
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_page_break()
    
    # Índice (TOC)
    toc_title = doc.add_paragraph("Índice")
    toc_title.runs[0].font.name = 'Arial'
    toc_title.runs[0].font.size = Pt(18)
    toc_title.runs[0].font.color.rgb = RGBColor(*NAVY_HEX)
    toc_title.runs[0].font.bold = True
    toc_title.runs[0].font.bold = True
    
    p_inst = doc.add_paragraph("(Haz clic derecho aquí y selecciona 'Actualizar campos' para generar la tabla de contenido)")
    p_inst.runs[0].font.name = 'Arial'
    p_inst.runs[0].font.size = Pt(9)
    p_inst.runs[0].font.italic = True
    p_inst.runs[0].font.color.rgb = RGBColor(120, 120, 120)
    
    add_toc(doc)
    
    doc.add_page_break()
    
    # 2. Secciones
    order = [
        ("Resumen Ejecutivo", "resumen_ejecutivo"),
        ("Alcance Funcional", "alcance_funcional"),
        ("Arquitectura Propuesta", "arquitectura"),
        ("Metodología", "metodologia"),
        ("Plan de Sprints", "plan_sprints"),
        ("Supuestos", "supuestos"),
        ("Exclusiones", "exclusiones"),
        ("Inversión", "inversion"),
    ]
    
    for title_text, key in order:
        if key in sections and sections[key]:
            # Encabezado de la sección
            h = doc.add_heading(title_text, level=1)
            h_run = h.runs[0]
            h_run.font.name = 'Arial'
            h_run.font.color.rgb = RGBColor(*TEAL_HEX)
            
            content = sections[key]
            
            # Procesamiento simple de párrafos, listas y TABLAS
            in_table = False
            current_table = None
            
            for line in content.split("\n"):
                line = line.strip()
                if not line:
                    in_table = False
                    continue
                
                if line.startswith("|") and line.endswith("|"):
                    if not in_table:
                        in_table = True
                        current_table = doc.add_table(rows=0, cols=0)
                        current_table.style = 'Table Grid'
                    
                    # Remove edge pipes and split
                    cells = [c.strip() for c in line.strip('|').split('|')]
                    
                    # Skip markdown separator line like |---|---|
                    if all(c.replace("-", "").strip() == "" for c in cells):
                        continue
                    
                    # Ensure table has enough columns
                    if len(current_table.columns) == 0:
                        for _ in cells:
                            current_table.add_column(Inches(1.2))
                            
                    row_cells = current_table.add_row().cells
                    for idx, cell_text in enumerate(cells):
                        if idx < len(row_cells):
                            row_cells[idx].text = cell_text
                else:
                    in_table = False
                    if line.startswith("- ") or line.startswith("* ") or line.startswith("• "):
                        clean_line = line[2:].strip().replace('**', '')
                        p = doc.add_paragraph(clean_line, style='List Bullet')
                    elif line.startswith("### "):
                        clean_line = line[4:].strip()
                        p = doc.add_heading(clean_line, level=3)
                        p.runs[0].font.color.rgb = RGBColor(*NAVY_HEX)
                    elif line.startswith("## "):
                        clean_line = line[3:].strip()
                        p = doc.add_heading(clean_line, level=2)
                        p.runs[0].font.color.rgb = RGBColor(*NAVY_HEX)
                    elif line.startswith("#"):
                        clean_line = line.lstrip("#").strip()
                        p = doc.add_heading(clean_line, level=2)
                        p.runs[0].font.color.rgb = RGBColor(*NAVY_HEX)
                    else:
                        p = doc.add_paragraph(line)
                    
            doc.add_page_break()
            
    # Añadir números de página
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("Página ")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(150, 150, 150)
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        r._r.append(fldChar1)
        r._r.append(instrText)
        r._r.append(fldChar2)

    # Crear directorio si no existe
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    
    doc.save(output_path)
    return output_path
